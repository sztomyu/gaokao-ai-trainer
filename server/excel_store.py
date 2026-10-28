# -*- coding: utf-8 -*-
"""
excel_store.py — 数据访问层
统一读写 5 个本地 Excel 文件（免数据库）：
  users.xlsx      账号/角色
  questions.xlsx  题目（主表 questions + 知识点树子表 knowledge_points）
  papers.xlsx     真题卷/模拟卷
  records.xlsx    作答记录/错题
  mastery.xlsx    知识点掌握度画像
所有写操作前自动备份为 .bak。
开发环境：数据存于 ./data/；打包安装后自动存于 %LOCALAPPDATA%/GaoKaoTrainer/data/。
"""
import os
import sys
import json
import shutil
import hashlib
import datetime
from openpyxl import Workbook, load_workbook

APP_NAME = "GaoKaoTrainer"
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
    DATA_DIR = os.path.join(
        os.environ.get("LOCALAPPDATA", os.path.expanduser("~")), APP_NAME, "data")
    SEED_SRC = os.path.join(sys._MEIPASS, "data")   # PyInstaller 只读模板
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    DATA_DIR = os.path.join(BASE_DIR, "data")
    SEED_SRC = DATA_DIR
os.makedirs(DATA_DIR, exist_ok=True)

F_USERS = os.path.join(DATA_DIR, "users.xlsx")
F_QUESTIONS = os.path.join(DATA_DIR, "questions.xlsx")
F_PAPERS = os.path.join(DATA_DIR, "papers.xlsx")
F_RECORDS = os.path.join(DATA_DIR, "records.xlsx")
F_MASTERY = os.path.join(DATA_DIR, "mastery.xlsx")

USERS_COLS = ["username", "password_hash", "role", "name", "status"]
QUESTION_COLS = ["qid", "subject", "chapter", "knowledge_point", "qtype",
                 "difficulty", "stem", "options", "answer", "analysis",
                 "source", "tag", "media_type", "media_paths", "created_at"]
KP_COLS = ["subject", "chapter", "knowledge_point"]
PAPER_COLS = ["pid", "title", "subject", "year", "ptype", "difficulty", "qids"]
RECORD_COLS = ["rid", "username", "qid", "user_answer", "correct", "duration_sec", "ts"]
MASTERY_COLS = ["username", "knowledge_point", "mastery", "status", "last_review", "review_count"]

QTYPES = ["单选", "多选", "填空", "解答"]
SUBJECTS = ["语文", "数学", "英语", "物理", "化学", "生物"]


def media_paths_of(q):
    """读取题目的图片清单（media_paths 存的是 JSON 字符串或 list）。"""
    v = q.get("media_paths", "[]")
    if isinstance(v, list):
        return v
    try:
        r = json.loads(v or "[]")
        return r if isinstance(r, list) else []
    except Exception:
        return []


def media_dir():
    """题目图形/图片统一存放目录：data/media/。"""
    d = os.path.join(DATA_DIR, "media")
    os.makedirs(d, exist_ok=True)
    return d


# ---------------- 基础读写 ----------------
def hash_pwd(pwd: str) -> str:
    return hashlib.sha256(("gk#" + pwd).encode("utf-8")).hexdigest()


def _read_sheet(path, sheet=None):
    if not os.path.exists(path):
        return []
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None for v in r):
            continue
        d = {}
        for i, col in enumerate(header):
            v = r[i] if i < len(r) else None
            d[col] = "" if v is None else v
        out.append(d)
    return out


def _write_sheet(path, cols, rows, sheet=None):
    """整表重写（数据量为题库级，性能足够）。写前备份 .bak。"""
    if os.path.exists(path):
        shutil.copy2(path, path + ".bak")
        wb = load_workbook(path)
    else:
        wb = Workbook()
    if sheet:
        if sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            ws = wb.create_sheet(sheet)
    else:
        ws = wb.active
    ws.delete_rows(1, ws.max_row if ws.max_row else 1)
    ws.append(cols)
    for row in rows:
        ws.append([row.get(c, "") for c in cols])
    # 新建工作簿时去掉多余的默认空 Sheet
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    wb.save(path)


def _next_id(rows, key, prefix):
    mx = 0
    for r in rows:
        v = str(r.get(key, ""))
        if v.startswith(prefix):
            try:
                mx = max(mx, int(v[len(prefix):]))
            except ValueError:
                pass
    return "%s%04d" % (prefix, mx + 1)


# ---------------- 初始化与种子数据 ----------------
def _seed_questions():
    """内置示例试题：数学/物理/英语，覆盖常见知识点，可直接训练，也可作为录入格式范例。"""
    def q(subject, chapter, kp, qtype, diff, stem, options, answer, analysis):
        return {"subject": subject, "chapter": chapter, "knowledge_point": kp,
                "qtype": qtype, "difficulty": diff, "stem": stem,
                "options": json.dumps(options, ensure_ascii=False) if options else "",
                "answer": answer, "analysis": analysis,
                "source": "系统内置示例", "created_at": str(datetime.date.today())}
    items = [
        # ---- 数学 ----
        q("数学", "函数与导数", "函数的概念与性质", "单选", 2,
          "函数 f(x)=√(x-1)+1/(x-3) 的定义域是（ ）",
          ["[1,+∞)", "(1,3)∪(3,+∞)", "[1,3)∪(3,+∞)", "(3,+∞)"],
          "C",
          "被开方数需 x-1≥0 即 x≥1；分母 x-3≠0 即 x≠3。取交集得 [1,3)∪(3,+∞)，选 C。"),
        q("数学", "函数与导数", "导数的几何意义", "单选", 3,
          "曲线 y=x³-2x 在点 (1,-1) 处的切线斜率为（ ）",
          ["-2", "1", "3", "5"],
          "B",
          "y′=3x²-2，代入 x=1 得 3-2=1，选 B。导数的几何意义即切线斜率。"),
        q("数学", "函数与导数", "利用导数研究单调性", "填空", 3,
          "函数 f(x)=x³-3x 的单调递增区间是________。",
          None,
          "(-∞,-1)和(1,+∞)",
          "f′(x)=3x²-3=3(x-1)(x+1)>0，解得 x<-1 或 x>1，故增区间为 (-∞,-1) 与 (1,+∞)。"),
        q("数学", "三角函数", "三角函数的图象与性质", "单选", 3,
          "函数 y=2sin(2x+π/3) 的最小正周期是（ ）",
          ["π/2", "π", "2π", "4π"],
          "B",
          "周期 T=2π/|ω|=2π/2=π，选 B。"),
        q("数学", "三角函数", "同角三角函数基本关系", "填空", 2,
          "已知 sinα=3/5，α为第二象限角，则 cosα=________。",
          None,
          "-4/5",
          "sin²α+cos²α=1 → cosα=±4/5；第二象限 cosα<0，故 cosα=-4/5。"),
        q("数学", "数列", "等差数列", "单选", 2,
          "等差数列 {aₙ} 中，a₁=2，d=3，则 a₁₀=（ ）",
          ["27", "29", "30", "32"],
          "B",
          "aₙ=a₁+(n-1)d=2+9×3=29，选 B。"),
        q("数学", "数列", "等比数列", "填空", 3,
          "等比数列 {aₙ} 中，a₁=1，q=2，则前 5 项和 S₅=________。",
          None,
          "31",
          "S₅=a₁(q⁵-1)/(q-1)=(32-1)/(2-1)=31。"),
        q("数学", "立体几何", "空间直线与平面的位置关系", "单选", 3,
          "设 m、n 是两条不同的直线，α 是一个平面，下列命题正确的是（ ）",
          ["若 m∥α，n∥α，则 m∥n", "若 m⊥α，n⊂α，则 m⊥n",
           "若 m∥n，n⊂α，则 m∥α", "若 m⊥n，n⊂α，则 m⊥α"],
          "B",
          "线面垂直的定义：m⊥α 则 m 垂直于 α 内任意直线，故 B 正确。A 中 m、n 可相交或异面；C 中 m 可能在 α 内；D 中 m 与 α 关系不定。"),
        q("数学", "概率统计", "古典概型", "单选", 2,
          "从 1、2、3、4、5 中任取两个不同的数，取到两个偶数的概率是（ ）",
          ["1/10", "1/5", "3/10", "2/5"],
          "A",
          "总取法 C(5,2)=10；偶数只有 2、4，取两个偶数仅 1 种，P=1/10，选 A。"),
        q("数学", "平面向量", "平面向量的数量积", "单选", 3,
          "已知 |a|=2，|b|=3，a 与 b 的夹角为 60°，则 a·b=（ ）",
          ["3", "3√3", "6", "6√3"],
          "A",
          "a·b=|a||b|cosθ=2×3×cos60°=6×1/2=3，选 A。"),
        q("数学", "解析几何", "椭圆的标准方程与性质", "单选", 4,
          "椭圆 x²/25+y²/16=1 的焦距为（ ）",
          ["3", "6", "8", "10"],
          "B",
          "c²=a²-b²=25-16=9，c=3，焦距 2c=6，选 B。"),
        q("数学", "不等式", "基本不等式", "单选", 4,
          "已知 x>0，y>0 且 x+y=4，则 xy 的最大值为（ ）",
          ["2", "4", "8", "16"],
          "B",
          "由基本不等式 xy≤((x+y)/2)²=4，当且仅当 x=y=2 时取等号，选 B。"),
        # ---- 物理 ----
        q("物理", "直线运动", "匀变速直线运动的规律", "单选", 2,
          "汽车以 20m/s 的速度做匀减速直线运动，加速度大小为 5m/s²，刹车后 3s 内的位移是（ ）",
          ["37.5m", "40m", "45m", "60m"],
          "A",
          "x=v₀t-½at²=20×3-½×5×9=60-22.5=37.5m。注意刹车 4s 停止，3s 未停，选 A。"),
        q("物理", "相互作用与牛顿运动定律", "牛顿第二定律", "单选", 3,
          "质量 2kg 的物体受到 10N 的水平拉力，在粗糙水平面上运动，动摩擦因数 μ=0.2（g 取 10m/s²），物体的加速度为（ ）",
          ["3m/s²", "4m/s²", "5m/s²", "6m/s²"],
          "A",
          "f=μmg=0.2×2×10=4N；F-f=ma → a=(10-4)/2=3m/s²，选 A。"),
        q("物理", "机械能", "机械能守恒定律", "填空", 3,
          "质量 1kg 的小球从 5m 高处自由落下（g 取 10m/s²），落地瞬间动能为________J。",
          None,
          "50",
          "机械能守恒：Eₖ=mgh=1×10×5=50J。"),
        q("物理", "静电场", "电场强度", "单选", 3,
          "点电荷 Q=2×10⁻⁶C 在真空中距其 0.3m 处产生的场强大小为（k=9×10⁹ N·m²/C²）（ ）",
          ["1×10⁵ N/C", "2×10⁵ N/C", "3×10⁵ N/C", "6×10⁵ N/C"],
          "B",
          "E=kQ/r²=9×10⁹×2×10⁻⁶/0.09=2×10⁵ N/C，选 B。"),
        q("物理", "恒定电流", "闭合电路欧姆定律", "单选", 3,
          "电源电动势 E=6V，内阻 r=1Ω，外接电阻 R=5Ω，则路端电压为（ ）",
          ["1V", "5V", "6V", "5.5V"],
          "B",
          "I=E/(R+r)=6/6=1A，U=IR=5V，选 B。"),
        q("物理", "电磁感应", "法拉第电磁感应定律", "单选", 4,
          "穿过单匝线圈的磁通量在 0.2s 内由 0.04Wb 均匀增加到 0.1Wb，则线圈中的感应电动势为（ ）",
          ["0.1V", "0.2V", "0.3V", "0.6V"],
          "C",
          "E=n·ΔΦ/Δt=1×(0.1-0.04)/0.2=0.3V，选 C。"),
        # ---- 英语 ----
        q("英语", "语法", "定语从句", "单选", 2,
          "This is the school ______ I studied three years ago.",
          ["which", "where", "that", "when"],
          "B",
          "先行词 school 表地点，从句中不缺主宾，用关系副词 where（=in which），选 B。"),
        q("英语", "语法", "虚拟语气", "单选", 4,
          "If I ______ time last night, I would have gone to the concert.",
          ["had", "have had", "had had", "would have"],
          "C",
          "与过去事实相反的虚拟语气，if 从句用过去完成时 had had，主句 would have done，选 C。"),
        q("英语", "语法", "时态与语态", "单选", 2,
          "By the end of last month, we ______ 2,000 English words.",
          ["learned", "have learned", "had learned", "learn"],
          "C",
          "by the end of+过去时间，主句用过去完成时 had learned，选 C。"),
        q("英语", "语法", "非谓语动词", "单选", 3,
          "______ from the top of the hill, the city looks beautiful.",
          ["Seeing", "Seen", "To see", "See"],
          "B",
          "主语 the city 与 see 是被动关系，用过去分词 Seen 作状语，选 B。"),
        q("英语", "阅读", "主旨大意", "单选", 3,
          "阅读主旨题解题的关键是（ ）",
          ["逐词翻译全文", "关注首尾段与各段主题句", "只看生词", "先做细节题"],
          "B",
          "主旨大意题应优先读首尾段与各段首句（主题句），把握文章结构与作者意图，选 B。"),
        q("英语", "写作", "书面表达句型", "填空", 2,
          "用英语表达“毫无疑问”，常用句型：There is no doubt ______...（填一个连词）",
          None,
          "that",
          "There is no doubt that... 为固定句型，that 引导同位语从句。"),
    ]
    rows = []
    for i, it in enumerate(items, 1):
        it = dict(it)
        it["qid"] = "Q%04d" % i
        rows.append(it)
    return rows


def _seed_kp():
    pairs = [
        ("数学", "函数与导数", "函数的概念与性质"), ("数学", "函数与导数", "导数的几何意义"),
        ("数学", "函数与导数", "利用导数研究单调性"), ("数学", "三角函数", "三角函数的图象与性质"),
        ("数学", "三角函数", "同角三角函数基本关系"), ("数学", "数列", "等差数列"),
        ("数学", "数列", "等比数列"), ("数学", "立体几何", "空间直线与平面的位置关系"),
        ("数学", "概率统计", "古典概型"), ("数学", "平面向量", "平面向量的数量积"),
        ("数学", "解析几何", "椭圆的标准方程与性质"), ("数学", "不等式", "基本不等式"),
        ("物理", "直线运动", "匀变速直线运动的规律"), ("物理", "相互作用与牛顿运动定律", "牛顿第二定律"),
        ("物理", "机械能", "机械能守恒定律"), ("物理", "静电场", "电场强度"),
        ("物理", "恒定电流", "闭合电路欧姆定律"), ("物理", "电磁感应", "法拉第电磁感应定律"),
        ("英语", "语法", "定语从句"), ("英语", "语法", "虚拟语气"),
        ("英语", "语法", "时态与语态"), ("英语", "语法", "非谓语动词"),
        ("英语", "阅读", "主旨大意"), ("英语", "写作", "书面表达句型"),
    ]
    return [{"subject": s, "chapter": c, "knowledge_point": k} for s, c, k in pairs]


def _seed_papers(qrows):
    math_ids = [r["qid"] for r in qrows if r["subject"] == "数学"][:10]
    phys_ids = [r["qid"] for r in qrows if r["subject"] == "物理"][:6]
    return [
        {"pid": "P0001", "title": "2025年数学高考真题卷（示例）", "subject": "数学",
         "year": 2025, "ptype": "真题", "difficulty": 3, "qids": ",".join(math_ids)},
        {"pid": "P0002", "title": "2024年物理高考真题卷（示例）", "subject": "物理",
         "year": 2024, "ptype": "真题", "difficulty": 3, "qids": ",".join(phys_ids)},
    ]


def _migrate_seeds():
    """PyInstaller 打包后首次运行：从只读模板目录复制种子数据到可写数据目录。"""
    if not os.path.exists(SEED_SRC):
        return
    for fname in os.listdir(SEED_SRC):
        src = os.path.join(SEED_SRC, fname)
        dst = os.path.join(DATA_DIR, fname)
        if os.path.isfile(src) and not os.path.exists(dst) and fname.endswith(".xlsx"):
            shutil.copy2(src, dst)


def init_all(force=False):
    """首次运行初始化 5 个 Excel 文件（含示例数据）。force=True 会重置全部数据。"""
    if getattr(sys, 'frozen', False) and not force:
        _migrate_seeds()
    def need(path):
        return force or not os.path.exists(path)

    if need(F_USERS):
        rows = [
            {"username": "teacher", "password_hash": hash_pwd("123456"),
             "role": "teacher", "name": "王老师", "status": "启用"},
            {"username": "student", "password_hash": hash_pwd("123456"),
             "role": "student", "name": "李同学", "status": "启用"},
        ]
        _write_sheet(F_USERS, USERS_COLS, rows)

    if need(F_QUESTIONS):
        qrows = _seed_questions()
        _write_sheet(F_QUESTIONS, QUESTION_COLS, qrows, sheet="questions")
        _write_sheet(F_QUESTIONS, KP_COLS, _seed_kp(), sheet="knowledge_points")

    if need(F_PAPERS):
        _write_sheet(F_PAPERS, PAPER_COLS, _seed_papers(_read_sheet(F_QUESTIONS, "questions")))

    if need(F_RECORDS):
        _write_sheet(F_RECORDS, RECORD_COLS, [])

    if need(F_MASTERY):
        _write_sheet(F_MASTERY, MASTERY_COLS, [])


# ---------------- 用户 ----------------
def verify_login(username, password):
    for r in _read_sheet(F_USERS):
        if str(r["username"]) == username and r["status"] == "启用" \
                and r["password_hash"] == hash_pwd(password):
            return {"username": r["username"], "role": r["role"], "name": r["name"]}
    return None


def list_users():
    return _read_sheet(F_USERS)


def save_user(username, password, role, name, status="启用"):
    rows = _read_sheet(F_USERS)
    for r in rows:
        if str(r["username"]) == username:
            if password:
                r["password_hash"] = hash_pwd(password)
            r["role"], r["name"], r["status"] = role, name, status
            _write_sheet(F_USERS, USERS_COLS, rows)
            return True
    rows.append({"username": username, "password_hash": hash_pwd(password or "123456"),
                 "role": role, "name": name, "status": status})
    _write_sheet(F_USERS, USERS_COLS, rows)
    return True


# ---------------- 题目 ----------------
def load_questions(subject=None, chapter=None, kp=None, qtype=None, keyword=None, tag=None):
    rows = _read_sheet(F_QUESTIONS, "questions")
    def hit(r):
        if subject and r["subject"] != subject:
            return False
        if chapter and r["chapter"] != chapter:
            return False
        if kp and kp not in str(r["knowledge_point"]):
            return False
        if qtype and r["qtype"] != qtype:
            return False
        if tag and tag not in str(r.get("tag", "")):
            return False
        if keyword and keyword not in str(r["stem"]):
            return False
        return True
    return [r for r in rows if hit(r)]


def get_question(qid):
    for r in _read_sheet(F_QUESTIONS, "questions"):
        if str(r["qid"]) == str(qid):
            return r
    return None


def get_questions_by_ids(qids):
    wanted = {str(x) for x in qids}
    return [r for r in _read_sheet(F_QUESTIONS, "questions") if str(r["qid"]) in wanted]


def _norm_stem(s):
    """标准化题干文本用于查重：去多余空白、统一标点、截取前 200 字符。"""
    t = str(s).strip()
    t = "".join(t.split())                    # 合并所有空白
    for a, b in [("．", "."), ("：", ":"), ("，", ","), ("；", ";"),
                 ("“", "\""), ("”", "\""), ("‘", "'"), ("’", "'")]:
        t = t.replace(a, b)
    return t[:200].lower()


def add_questions(items):
    """批量入库，自动按题干去重。返回 (new_qids, skip_count)。"""
    rows = _read_sheet(F_QUESTIONS, "questions")
    # 现有题库的去重指纹
    existing = {_norm_stem(r.get("stem", "")) for r in rows}
    new_ids, skipped = [], 0
    for it in items:
        it = dict(it)
        stem = it.get("stem", "")
        if not stem.strip():
            skipped += 1
            continue
        ns = _norm_stem(stem)
        if ns in existing:
            skipped += 1
            continue
        existing.add(ns)                            # 同时也防批量内重复
        qid = _next_id(rows, "qid", "Q")
        it["qid"] = qid
        if isinstance(it.get("options"), list):
            it["options"] = json.dumps(it["options"], ensure_ascii=False)
        if isinstance(it.get("media_paths"), list):
            it["media_paths"] = json.dumps(it["media_paths"], ensure_ascii=False)
        it.setdefault("created_at", str(datetime.date.today()))
        it.setdefault("tag", "训练题")
        it.setdefault("media_type", "无")
        it.setdefault("media_paths", "[]")
        rows.append(it)
        new_ids.append(qid)
        # 自动登记新知识点
        if it.get("subject") and it.get("knowledge_point"):
            for kp in str(it["knowledge_point"]).split(","):
                kp = kp.strip()
                if kp:
                    add_knowledge_point(it["subject"], it.get("chapter", "") or "未分章", kp)
    _write_sheet(F_QUESTIONS, QUESTION_COLS, rows, sheet="questions")
    return new_ids, skipped


def update_question(qid, fields):
    rows = _read_sheet(F_QUESTIONS, "questions")
    for r in rows:
        if str(r["qid"]) == str(qid):
            if isinstance(fields.get("options"), list):
                fields["options"] = json.dumps(fields["options"], ensure_ascii=False)
            r.update(fields)
            _write_sheet(F_QUESTIONS, QUESTION_COLS, rows, sheet="questions")
            return True
    return False


def delete_question(qid):
    rows = [r for r in _read_sheet(F_QUESTIONS, "questions") if str(r["qid"]) != str(qid)]
    _write_sheet(F_QUESTIONS, QUESTION_COLS, rows, sheet="questions")


def parse_options(q):
    """把 options 字段解析为选项列表。"""
    raw = q.get("options", "")
    if not raw:
        return []
    try:
        v = json.loads(raw)
        return v if isinstance(v, list) else []
    except Exception:
        return [s for s in str(raw).split("\n") if s.strip()]


# ---------------- 知识点树 ----------------
def get_kp_tree():
    """{subject: {chapter: [kp,...]}}"""
    tree = {}
    for r in _read_sheet(F_QUESTIONS, "knowledge_points"):
        tree.setdefault(r["subject"], {}).setdefault(r["chapter"], [])
        if r["knowledge_point"] not in tree[r["subject"]][r["chapter"]]:
            tree[r["subject"]][r["chapter"]].append(r["knowledge_point"])
    return tree


def add_knowledge_point(subject, chapter, kp):
    rows = _read_sheet(F_QUESTIONS, "knowledge_points")
    for r in rows:
        if r["subject"] == subject and r["chapter"] == chapter and r["knowledge_point"] == kp:
            return
    rows.append({"subject": subject, "chapter": chapter, "knowledge_point": kp})
    _write_sheet(F_QUESTIONS, KP_COLS, rows, sheet="knowledge_points")


# ---------------- 试卷 ----------------
def list_papers(subject=None, year=None, ptype=None):
    rows = _read_sheet(F_PAPERS)
    out = []
    for r in rows:
        if subject and r["subject"] != subject:
            continue
        if year and str(r["year"]) != str(year):
            continue
        if ptype and r["ptype"] != ptype:
            continue
        out.append(r)
    return out


def add_paper(title, subject, year, ptype, difficulty, qids):
    rows = _read_sheet(F_PAPERS)
    pid = _next_id(rows, "pid", "P")
    rows.append({"pid": pid, "title": title, "subject": subject, "year": year,
                 "ptype": ptype, "difficulty": difficulty, "qids": ",".join(map(str, qids))})
    _write_sheet(F_PAPERS, PAPER_COLS, rows)
    return pid


# ---------------- 作答记录 ----------------
def add_record(username, qid, user_answer, correct, duration_sec=0):
    rows = _read_sheet(F_RECORDS)
    rid = _next_id(rows, "rid", "R")
    rows.append({"rid": rid, "username": username, "qid": qid,
                 "user_answer": user_answer, "correct": 1 if correct else 0,
                 "duration_sec": duration_sec,
                 "ts": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")})
    _write_sheet(F_RECORDS, RECORD_COLS, rows)
    return rid


def load_records(username=None, qid=None, correct=None):
    rows = _read_sheet(F_RECORDS)
    out = []
    for r in rows:
        if username and r["username"] != username:
            continue
        if qid and str(r["qid"]) != str(qid):
            continue
        if correct is not None and int(r["correct"] or 0) != (1 if correct else 0):
            continue
        out.append(r)
    return out


# ---------------- 掌握度 ----------------
def get_mastery(username, kp=None):
    rows = [r for r in _read_sheet(F_MASTERY) if r["username"] == username]
    if kp is not None:
        for r in rows:
            if r["knowledge_point"] == kp:
                return r
        return None
    return rows


def upsert_mastery(username, kp, mastery, status, last_review=None, review_count=0):
    rows = _read_sheet(F_MASTERY)
    for r in rows:
        if r["username"] == username and r["knowledge_point"] == kp:
            r["mastery"] = mastery
            r["status"] = status
            r["last_review"] = last_review if last_review is not None else r.get("last_review", "")
            r["review_count"] = review_count
            _write_sheet(F_MASTERY, MASTERY_COLS, rows)
            return
    rows.append({"username": username, "knowledge_point": kp, "mastery": mastery,
                 "status": status, "last_review": last_review or "",
                 "review_count": review_count})
    _write_sheet(F_MASTERY, MASTERY_COLS, rows)
