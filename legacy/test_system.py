# -*- coding: utf-8 -*-
"""
test_system.py — 模块级自动化验证（无界面，沙箱数据目录，不影响真实题库）
覆盖：Excel初始化/登录/判分/掌握度闭环/组卷/复测/学情/CRUD/
      识别演示模式/长文本分段识别/截断JSON抢救/试卷
运行：python test_system.py
"""
import os
import sys
import json
import tempfile
import datetime

# ---- 沙箱：把数据目录重定向到临时目录，避免污染真实题库 ----
_TMP = tempfile.mkdtemp(prefix="gk_test_")
import excel_store as store
store.DATA_DIR = _TMP
store.F_USERS = os.path.join(_TMP, "users.xlsx")
store.F_QUESTIONS = os.path.join(_TMP, "questions.xlsx")
store.F_PAPERS = os.path.join(_TMP, "papers.xlsx")
store.F_RECORDS = os.path.join(_TMP, "records.xlsx")
store.F_MASTERY = os.path.join(_TMP, "mastery.xlsx")

import engine
import study_loop
import analytics
import ai_recognize
ai_recognize.SETTINGS_PATH = os.path.join(_TMP, "settings.json")

PASS, FAIL = 0, 0


def check(name, cond):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  [PASS] %s" % name)
    else:
        FAIL += 1
        print("  [FAIL] %s" % name)


def main():
    print("== 1. 初始化与Excel文件 ==")
    store.init_all()
    for f in [store.F_USERS, store.F_QUESTIONS, store.F_PAPERS, store.F_RECORDS, store.F_MASTERY]:
        check("文件存在 %s" % os.path.basename(f), os.path.exists(f))
    qs = store.load_questions()
    check("内置示例题 >= 20 道", len(qs) >= 20)
    check("题目字段完整", all(k in qs[0] for k in
          ["qid", "subject", "chapter", "knowledge_point", "qtype", "difficulty",
           "stem", "options", "answer", "analysis"]))
    from openpyxl import load_workbook
    wb = load_workbook(store.F_QUESTIONS)
    check("questions.xlsx 含 questions/knowledge_points 两个子表且无多余空表",
          wb.sheetnames == ["questions", "knowledge_points"])
    wb.close()

    print("== 2. 登录认证 ==")
    check("教师登录", store.verify_login("teacher", "123456")["role"] == "teacher")
    check("学生登录", store.verify_login("student", "123456")["role"] == "student")
    check("错误密码拒绝", store.verify_login("student", "wrong") is None)

    print("== 2b. 双 API 配置(text + vision) 独立保存 ==")
    ai_recognize.save_settings(
        text={"api_base": "https://api.deepseek.com", "api_key": "text-key-abc", "model": "deepseek-v4-pro"},
        vision={"api_base": "https://dashscope.aliyuncs.com/compatible-mode",
                "api_key": "vision-key-xyz", "model": "qwen-vl-max"})
    cfg = ai_recognize.load_settings()
    check("text.api_key 独立保存", cfg["text"]["api_key"] == "text-key-abc")
    check("vision.api_key 独立保存", cfg["vision"]["api_key"] == "vision-key-xyz")
    check("vision.api_base 独立", cfg["vision"]["api_base"] == "https://dashscope.aliyuncs.com/compatible-mode")
    check("api_configured(text)", ai_recognize.api_configured("text") is True)
    check("api_configured(vision)", ai_recognize.api_configured("vision") is True)
    import json as _json, base64 as _b64
    legacy = {"api_base": "https://api.deepseek.com",
              "api_key_enc": _b64.b64encode(b"old-key").decode(),
              "model": "deepseek-chat", "vision_model": "deepseek-vl"}
    _json.dump(legacy, open(ai_recognize.SETTINGS_PATH, "w", encoding="utf-8"), ensure_ascii=False)
    cfg2 = ai_recognize.load_settings()
    check("旧版 text.api_key 迁移", cfg2["text"]["api_key"] == "old-key")
    check("旧版 vision.model 迁移", cfg2["vision"]["model"] == "deepseek-vl")

    print("== 3. 判分 ==")
    q_choice = next(q for q in qs if q["qtype"] == "单选")
    check("单选判对", engine.grade(q_choice, q_choice["answer"]) is True)
    check("单选判错", engine.grade(q_choice, "Z") is False)
    q_fill = next(q for q in qs if q["qtype"] == "填空")
    check("填空判对", engine.grade(q_fill, q_fill["answer"]) is True)
    check("解答题返回None(自评)", engine.grade({"qtype": "解答", "answer": "x"}, "y") is None)

    print("== 4. 作答->掌握度->相似题闭环 ==")
    kp = "古典概型"
    q1 = store.load_questions(kp=kp)[0]
    r1 = study_loop.on_answer("student", q1, q1["answer"])
    check("作答记录写入", len(store.load_records(username="student")) == 1)
    check("答对后掌握度上升", r1["mastery_after"] > r1["mastery_before"])
    r2 = study_loop.on_answer("student", q1, "Z")
    check("答错后掌握度下降", r2["mastery_after"] < r1["mastery_after"])
    sim = study_loop.similar_question("student", kp)
    check("相似题可取", sim is not None and kp in sim["knowledge_point"])

    print("== 5. 掌握度过关与复测调度 ==")
    for _ in range(8):
        study_loop.on_answer("student", q1, q1["answer"])
    m = store.get_mastery("student", kp)
    check("掌握度>=80判定过关", m["status"] == "已过关" and float(m["mastery"]) >= 80)
    check("无到期复测(刚过关)", len(engine.retest_due("student")) == 0)
    rows = store._read_sheet(store.F_MASTERY)
    for r in rows:
        if r["username"] == "student" and r["knowledge_point"] == kp:
            r["last_review"] = str(datetime.date.today() - datetime.timedelta(days=3))
    store._write_sheet(store.F_MASTERY, store.MASTERY_COLS, rows)
    due = engine.retest_due("student")
    check("3天后复测到期", len(due) == 1 and due[0]["knowledge_point"] == kp)
    study_loop.apply_retest_result("student", kp, True)
    check("复测通过review_count+1", store.get_mastery("student", kp)["review_count"] == 1)

    print("== 6. 组卷算法 ==")
    paper = engine.generate_paper("数学", 3, 8)
    check("生成8道题", len(paper) == 8)
    check("题目不重复", len({q["qid"] for q in paper}) == len(paper))
    pid = store.add_paper("单元测试卷", "数学", 2026, "模拟", 3, [q["qid"] for q in paper])
    check("试卷保存", any(p["pid"] == pid for p in store.list_papers()))

    print("== 7. 学情分析 ==")
    ov = analytics.overall("student")
    check("总览统计", ov["total_answers"] >= 10 and ov["kp_passed"] >= 1)
    summary = analytics.kp_summary("student")
    check("知识点画像含古典概型", any(s["kp"] == kp for s in summary))

    print("== 8. 教师端：题目CRUD与去重 ==")
    new_ids, sk0 = store.add_questions([{
        "subject": "数学", "chapter": "数列", "knowledge_point": "等差数列",
        "qtype": "单选", "difficulty": 2, "stem": "测试新增题：1+1=?",
        "options": ["1", "2", "3", "4"], "answer": "B", "analysis": "1+1=2",
        "source": "自动化测试"}])
    check("新增题目", store.get_question(new_ids[0]) is not None)
    # 重复题应被跳过
    dup_ids, dup_sk = store.add_questions([{
        "stem": "测试新增题：1+1=?", "subject": "数学", "chapter": "数列",
        "qtype": "单选", "difficulty": 2, "options": ["1", "2"], "answer": "B"}])
    check("重复题干被跳过", len(dup_ids) == 0 and dup_sk == 1)
    # 批量内重复也跳过
    batch_ids, batch_sk = store.add_questions([
        {"stem": "批量重复题", "subject": "物理", "qtype": "单选"},
        {"stem": "批量重复题", "subject": "物理", "qtype": "单选"}])
    check("批量内重复跳过", len(batch_ids) == 1 and batch_sk == 1)
    store.update_question(new_ids[0], {"difficulty": 4})
    check("修改题目", int(store.get_question(new_ids[0])["difficulty"]) == 4)
    store.delete_question(new_ids[0])
    store.delete_question(batch_ids[0])   # 清理
    check("删除题目", store.get_question(new_ids[0]) is None)

    print("== 9. 长文本分段（不切断题目） ==")
    filler = "请写出详细的计算过程，并对结果进行检验与讨论。" * 3
    text40 = "2026数学卷\n" + "\n".join(
        "%d. 第%d题题干：计算 %d+%d=？%s\nA. %d  B. %d  C. %d  D. %d"
        % (i, i, i, i, filler, 2 * i, 2 * i + 1, 2 * i + 2, 2 * i + 3)
        for i in range(1, 41))
    prefix, blocks = ai_recognize._split_question_blocks(text40)
    check("切出40个完整题目块", len(blocks) == 40)
    check("卷头提取", "2026数学卷" in prefix)
    chunks = ai_recognize._chunk_blocks(blocks, max_chars=2500)
    check("分段数>1", len(chunks) > 1)
    check("每段<=2500字", all(len(c) <= 2600 for c in chunks))  # 单块不切断故留余量
    check("分段不丢题", sum(c.count("\n\n") + 1 for c in chunks) == 40)
    check("每段以题号开头", all(c.lstrip()[0].isdigit() for c in chunks))

    print("== 10. 截断JSON抢救解析 ==")
    good = '[{"stem":"题1","answer":"A"},{"stem":"题2","options":["x","y"],"answer":"B"}]'
    check("完整数组正常解析", len(ai_recognize._parse_json_array(good)) == 2)
    truncated = ('```json\n[{"stem":"题1","answer":"A"},'
                 '{"stem":"题2","options":["x","y"],"answer":"B"},'
                 '{"stem":"题3被截断","opt')
    salvaged = ai_recognize._parse_json_array(truncated)
    check("截断输出抢救回2道完整题", len(salvaged) == 2 and salvaged[1]["answer"] == "B")
    try:
        ai_recognize._parse_json_array("抱歉，我无法处理这张图片。")
        check("无JSON时抛错", False)
    except ValueError:
        check("无JSON时抛错", True)

    print("== 10b. 畸形JSON宽松修复（数学内容常见破坏） ==")
    trailing = '[{"stem":"题1","answer":"A",},]'
    check("尾逗号修复", len(ai_recognize._parse_json_array(trailing)) == 1)
    newline_in_str = '[{"stem":"第一行\n第二行","answer":"A"}]'.replace("\\n", "\n")
    got = ai_recognize._parse_json_array(newline_in_str)
    check("字符串内真实换行修复", len(got) == 1 and "第一行" in got[0]["stem"])
    latex = '[{"stem":"求 \\( x^2+1 \\) 在 x=1 处的值","answer":"2"}]'
    got2 = ai_recognize._parse_json_array(latex)
    check("LaTeX反斜杠修复", len(got2) == 1 and "x^2" in got2[0]["stem"]
          and "\\(" in got2[0]["stem"])

    print("== 11. 分段识别合并（模拟API） ==")
    calls = {"n": 0}
    def fake_chat(messages, kind="text", model=None, max_tokens=8000):
        calls["n"] += 1
        user = messages[-1]["content"]
        if calls["n"] in (2, 3):              # 模拟某段持续失败
            return "这不是JSON"
        # 适配新切题格式：块前缀为 "X题 " 而非 "X. "
        pairs = __import__("re").findall(r"(\d+)\. |(\d+)题 ", user)
        nums = [int(a or b) for a, b in pairs]
        return json.dumps([{"subject": "数学", "chapter": "测试", "knowledge_point": "分段识别",
                            "qtype": "单选", "difficulty": 2,
                            "stem": "第%d题" % n, "options": ["a", "b"], "answer": "A",
                            "analysis": "x", "confidence": 0.9} for n in nums],
                          ensure_ascii=False)
    orig_chat = ai_recognize._chat
    ai_recognize._chat = fake_chat
    try:
        data, n_seg, n_fail, last_err = ai_recognize._recognize_text(text40, "数学")
        check("全部分段被调用", calls["n"] >= n_seg)
        check("失败段被跳过、其余题合并", len(data) + n_fail * (40 // n_seg) >= 30)
        check("失败段计数正确", n_fail >= 1)
        check("保留最后一次错误信息", bool(last_err))
    finally:
        ai_recognize._chat = orig_chat

    print("== 12. 大模型识别（演示模式，未配置API） ==")
    # 12 段依赖"未配置 API"路径：清空 2b 测试写入的 key
    ai_recognize.save_settings(
        text={"api_base": "https://api.deepseek.com", "api_key": "", "model": "deepseek-v4-pro"},
        vision={"api_base": "", "api_key": "", "model": ""})
    demo_txt = os.path.join(_TMP, "demo_paper.txt")
    with open(demo_txt, "w", encoding="utf-8") as f:
        f.write("2026年数学模拟卷\n"
                "1. 已知集合A={1,2,3}，B={2,3,4}，则A∩B=（ ）\n"
                "A. {1}  B. {2,3}  C. {4}  D. 空集\n"
                "2. 函数y=x²的最小值是（ ）\n"
                "A. 0  B. 1  C. -1  D. 2\n"
                "3. 等差数列中a1=1,d=2，则a5=____。\n")
    items, msg = ai_recognize.recognize(demo_txt)
    check("演示模式切分出3道题", len(items) == 3)
    check("识别出选择题选项", len(items[0]["options"]) >= 2)
    check("演示模式低置信度待校对", all(it["confidence"] < 0.7 for it in items))
    check("演示模式tag为训练题", all(it.get("tag") == "训练题" for it in items))
    ids2, _ = store.add_questions(items)
    check("识别结果可入库", len(ids2) == 3 and store.get_question(ids2[0]) is not None)

    print("== 12b. 标签字段（tag）独立验证 ==")
    # add_questions 时缺少 tag 字段应自动补"训练题"
    new_ids, _ = store.add_questions([{
        "subject": "物理", "chapter": "力学", "knowledge_point": "牛顿定律",
        "qtype": "单选", "difficulty": 2, "stem": "无tag字段题目",
        "options": [], "answer": "A", "analysis": "", "source": "测试",
    }])
    check("无tag字段自动补'训练题'", (store.get_question(new_ids[0]) or {}).get("tag") == "训练题")
    # AI 提取的 tag 形式为 年/地区/类型
    nrm = ai_recognize._normalize_items(
        [{"stem": "2023年北京卷高考真题", "options": [], "answer": "A",
          "analysis": "2023年北京高考", "subject": "数学"}],
        "mock.docx")
    check("tag含2023/北京",
          "2023" in nrm[0]["tag"] and "北京" in nrm[0]["tag"])
    nrm2 = ai_recognize._normalize_items(
        [{"stem": "某通用练习", "options": [], "answer": "A", "analysis": "",
          "subject": "数学"}], "no source")
    check("无法识别时tag=训练题", nrm2[0]["tag"] == "训练题")
    # 标签过滤
    listed = store.load_questions(tag="训练题")
    check("标签过滤生效(至少1条)", len(listed) >= 1)

    print("== 12c. 看板统计（各科×题型分类） ==")
    stats = analytics.category_stats()
    check("category_stats 返回 total 字段", "total" in stats and stats["total"] >= 20)
    check("by_subject 是 list", isinstance(stats["by_subject"], list))
    check("by_category 至少 4 个分类", len(stats["by_category"]) >= 4)
    by_cat = stats["by_category"]
    check("分类为 4 项 (选择/填空/阅读/大题)",
          set(by_cat.keys()) >= {"选择题", "填空题", "阅读题", "大题"})
    check("category_of 单选->选择题", analytics.category_of({"qtype": "单选"}) == "选择题")
    check("category_of 多选->选择题", analytics.category_of({"qtype": "多选"}) == "选择题")
    check("category_of 填空->填空题", analytics.category_of({"qtype": "填空"}) == "填空题")
    check("category_of 解答(阅读)->阅读题",
          analytics.category_of({"qtype": "解答", "knowledge_point": "现代文阅读"}) == "阅读题")
    check("category_of 解答(其他)->大题",
          analytics.category_of({"qtype": "解答", "knowledge_point": "数列求和"}) == "大题")

    print("== 13. Excel备份机制 ==")
    check("写入生成.bak备份", os.path.exists(store.F_QUESTIONS + ".bak"))

    print("\n===== 结果：%d 通过，%d 失败 =====" % (PASS, FAIL))
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
