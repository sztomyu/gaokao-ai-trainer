# -*- coding: utf-8 -*-
"""
gen_knowledge_map.py — 考点提取手册生成器
从题库读取指定科目，把碎片章节归并为标准模块，提取考点并统计频次，
生成格式整齐的 Word 考点手册（表格化）。

题号与《<科目>高考练习.docx》习题集全局编号一致，可对照练习。

用法：
  python gen_knowledge_map.py --subject 数学 --output "数学高考考点整理.docx"
"""
import argparse
import datetime
import os
import re
import sys
from collections import Counter, OrderedDict

import openpyxl
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_exercise_book import latex_to_unicode, set_cn_font, add_page_number

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "server", "data", "questions.xlsx")

# ---------------- 章节归并为大模块 ----------------
# 规则按顺序匹配：更具体的关键词必须先于泛化关键词（如"三角"先于"函数"，
# 否则「三角函数」会被"函数"抢走）；兜底「算法与其他选讲」
_MODULE_RULES = [
    ("三角函数与解三角形", ["三角"]),
    ("平面向量", ["向量"]),
    ("数列", ["数列"]),
    ("不等式与线性规划", ["不等式", "线性规划"]),
    ("立体几何", ["立体几何", "空间几何", "三视图"]),
    ("平面解析几何", ["解析几何", "圆锥曲线", "双曲线", "抛物线", "椭圆"]),
    ("概率与统计", ["概率", "统计", "随机"]),
    ("计数原理与排列组合", ["计数", "排列", "组合"]),
    ("算法与其他选讲", ["算法", "参数方程", "极坐标", "选讲"]),
    ("集合与常用逻辑用语", ["集合", "逻辑"]),
    ("复数", ["复数"]),
    ("函数与导数", ["函数", "导数", "极限"]),
]
_FALLBACK_MODULE = "算法与其他选讲"


def chapter_to_module(ch):
    for name, kws in _MODULE_RULES:
        if any(k in ch for k in kws):
            return name
    return _FALLBACK_MODULE


def split_kps(kp_str):
    out = []
    for k in re.split(r"[,，、;；]", kp_str or ""):
        k = k.strip()
        if k:
            out.append(k)
    return out


def para(doc, text="", size=10.5, bold=False, color=None, align=None,
         space_after=4, cn="宋体"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    set_cn_font(r, cn=cn)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p, r


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_text(cell, text, size=10, bold=False, color=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    set_cn_font(r)
    p.paragraph_format.space_after = Pt(1)


def make_table(doc, header, rows, widths=None):
    """格式整齐的表格：表头深蓝底白字，斑马纹，网格线。"""
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j, h in enumerate(header):
        _cell_text(hdr[j], h, size=10, bold=True, color=(0xFF, 0xFF, 0xFF), center=True)
        _shade_cell(hdr[j], "2F5496")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            _cell_text(cells[j], str(v), size=10,
                       center=(j != 0 and len(str(v)) < 20))
            if i % 2 == 1:
                _shade_cell(cells[j], "F2F5FA")
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


def build(subject, output):
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["questions"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header)}
    items = []
    for r in rows[1:]:
        if not r or not r[idx.get("qid", 0)]:
            continue
        if str(r[idx["subject"]] or "").strip() != subject:
            continue
        items.append({
            "qid": str(r[idx["qid"]]),
            "chapter": str(r[idx["chapter"]] or "未分章").strip() or "未分章",
            "kps": split_kps(str(r[idx["knowledge_point"]] or "")),
            "stem": str(r[idx["stem"]] or "").strip(),
            "qtype": str(r[idx["qtype"]] or "").strip(),
        })
    wb.close()
    if not items:
        raise SystemExit("题库中没有科目为「%s」的题目" % subject)

    # 与习题集一致的全局题号：按 chapter 首次出现顺序遍历
    chapters = OrderedDict()
    for it in items:
        chapters.setdefault(it["chapter"], []).append(it)
    global_no = {}
    n = 0
    for ch, ch_items in chapters.items():
        for it in ch_items:
            n += 1
            global_no[id(it)] = n

    # 归并模块 → 章节 → 考点统计
    modules = OrderedDict()          # 模块名 -> Counter(考点 -> 频次)
    module_chapters = OrderedDict()  # 模块名 -> [章节名]
    module_qids = OrderedDict()      # 模块名 -> [题号]
    kp_qnos = {}                     # (模块, 考点) -> [题号]
    for ch, ch_items in chapters.items():
        mod = chapter_to_module(ch)
        modules.setdefault(mod, Counter())
        module_chapters.setdefault(mod, [])
        module_qids.setdefault(mod, [])
        if ch not in module_chapters[mod]:
            module_chapters[mod].append(ch)
        for it in ch_items:
            no = global_no[id(it)]
            module_qids[mod].append(no)
            for k in it["kps"]:
                modules[mod][k] += 1
                kp_qnos.setdefault((mod, k), []).append(no)

    total_q = len(items)
    total_kp = sum(len(c) for c in modules.values())

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_p)
    head_p = sec.header.paragraphs[0]
    head_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = head_p.add_run("高考%s 考点提取手册" % subject)
    hr.font.size = Pt(9)
    hr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    set_cn_font(hr)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    set_cn_font(normal)

    # ---------- 封面 ----------
    for _ in range(6):
        para(doc, "")
    para(doc, "高考%s 考点提取手册" % subject, size=30, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, cn="黑体", space_after=10)
    para(doc, "—— 从题库 %d 题提取 · %d 大模块 · %d 个考点 ——" % (total_q, len(modules), total_kp),
         size=13, color=(0x66, 0x66, 0x66), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    for ln in ["科目：%s" % subject,
               "总题量：%d 题（对应《%s高考练习.docx》习题集题号）" % (total_q, subject),
               "模块划分：%d 个标准模块（原碎片章节已归并）" % len(modules),
               "考点总数：%d 个（按题库知识点标签拆分统计）" % total_kp,
               "生成日期：%s" % datetime.date.today().strftime("%Y 年 %m 月 %d 日")]:
        para(doc, ln, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    doc.add_page_break()

    # ---------- 总览 ----------
    h = doc.add_heading("", level=1)
    r = h.add_run("一、模块总览")
    r.font.size = Pt(15); r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x66)
    set_cn_font(r, cn="黑体")
    rows_ov = []
    for i, (mod, cnt) in enumerate(modules.items(), 1):
        rows_ov.append(["第 %d 章　%s" % (i, mod),
                        "%d" % len(module_qids[mod]),
                        "%d" % len(cnt),
                        "；".join(module_chapters[mod])])
    make_table(doc, ["模块", "题数", "考点数", "归并的原章节"], rows_ov,
               widths=[5.5, 1.6, 1.6, 8.3])
    para(doc, "", space_after=2)
    para(doc, "注：「归并的原章节」为题库识别产生的碎片章节名，已按关键词归入左侧标准模块。",
         size=9, color=(0x88, 0x88, 0x88))

    # ---------- 高频考点 TOP20 ----------
    h = doc.add_heading("", level=1)
    r = h.add_run("二、高频考点 TOP 20")
    r.font.size = Pt(15); r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x66)
    set_cn_font(r, cn="黑体")
    para(doc, "按出现频次排序；同一考点跨模块出现时按各自模块单独统计。",
         size=9, color=(0x88, 0x88, 0x88), space_after=6)
    flat = []
    for mod, cnt in modules.items():
        for k, v in cnt.items():
            flat.append((k, v, mod))
    flat.sort(key=lambda x: (-x[1], x[0]))
    rows_top = [["%d" % (i + 1), k, "%d" % v, mod]
                for i, (k, v, mod) in enumerate(flat[:20])]
    make_table(doc, ["排名", "考点", "频次", "所属模块"], rows_top,
               widths=[1.4, 7.0, 1.6, 7.0])
    doc.add_page_break()

    # ---------- 各模块明细 ----------
    h = doc.add_heading("", level=1)
    r = h.add_run("三、分模块考点明细")
    r.font.size = Pt(15); r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x66)
    set_cn_font(r, cn="黑体")
    para(doc, "题号列与《%s高考练习.docx》习题集编号一致，可直接对照练习。" % subject,
         size=9, color=(0x88, 0x88, 0x88), space_after=8)

    for i, (mod, cnt) in enumerate(modules.items(), 1):
        h2 = doc.add_heading("", level=2)
        r2 = h2.add_run("第 %d 章　%s（%d 题 / %d 个考点）"
                        % (i, mod, len(module_qids[mod]), len(cnt)))
        r2.font.size = Pt(13); r2.font.bold = True
        r2.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
        set_cn_font(r2, cn="黑体")

        detail = []
        for rank, (k, v) in enumerate(sorted(cnt.items(), key=lambda x: (-x[1], x[0])), 1):
            nos = sorted(kp_qnos.get((mod, k), []))
            detail.append(["%d" % rank, k, "%d" % v, "、".join(str(x) for x in nos)])
        make_table(doc, ["序号", "考点", "频次", "对应题号"], detail,
                   widths=[1.3, 6.2, 1.4, 8.1])
        para(doc, "", space_after=4)

    doc.save(output)
    return total_q, len(modules), total_kp


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--subject", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()
    q, m, k = build(args.subject, args.output)
    print("OK：%s → %s（%d 题 / %d 模块 / %d 考点）" % (args.subject, args.output, q, m, k))


if __name__ == "__main__":
    main()
